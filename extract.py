import docx

def extract_text_from_docx(file_path):
    doc = docx.Document(file_path)
    full_text = []
    
    # Iterate through all elements in body in order is harder in python-docx, 
    # but we can just grab paragraphs then tables, or use a simpler approach.
    for para in doc.paragraphs:
        full_text.append(para.text)
        
    for table in doc.tables:
        for row in table.rows:
            row_data = []
            for cell in row.cells:
                row_data.append(cell.text.replace('\n', ' '))
            full_text.append(' | '.join(row_data))
            
    return '\n'.join(full_text)

text = extract_text_from_docx("Hermetik Astroloji.docx")
with open("sabian_symbols_full.txt", "w", encoding="utf-8") as f:
    f.write(text)
